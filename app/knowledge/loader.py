from __future__ import annotations

import csv
from pathlib import Path


class UnsupportedFileTypeError(ValueError):
    pass


def load_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        return path.read_text(errors="replace")
    if suffix == ".csv":
        return _load_csv_as_text(path)
    if suffix == ".pdf":
        return _load_pdf_as_text(path)
    if suffix == ".docx":
        return _load_docx_as_text(path)
    raise UnsupportedFileTypeError(f"Cannot load text from file type: {suffix}")


def _load_csv_as_text(path: Path) -> str:
    lines = []
    with path.open(newline="", errors="replace") as f:
        reader = csv.reader(f)
        for i, row in enumerate(reader):
            lines.append(", ".join(row))
            if i > 5000:
                lines.append("... (truncated)")
                break
    return "\n".join(lines)


def _load_pdf_as_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _load_docx_as_text(path: Path) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                parts.append(row_text)
    return "\n".join(parts)